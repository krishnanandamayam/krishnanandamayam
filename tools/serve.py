"""Local server: the reader at /, the HK editor at /edit. Never deployed.

    uv run python -m tools.serve [--port 8000] [--host 127.0.0.1]

The editor writes data/grantha.yaml and rebuilds site/. It never runs a git command that
changes anything: committing and pushing stay with you.
"""

from __future__ import annotations

import argparse
import json
import mimetypes
import re
import subprocess
import threading
import tomllib
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from io import StringIO
from urllib.parse import unquote, urlparse

from tools import build, inline, render, store
from tools import hk as HK

ROOT = store.ROOT
SITE = ROOT / "site"
EDITOR = ROOT / "tools" / "editor"
FLAGS = ROOT / "data" / "flags.toml"
LINT_OK = ROOT / "data" / "lint-ok.toml"
LOCK = threading.Lock()  # one writer / one Aksharamukha call at a time
FIELDS = ("title", "hk", "en", "te", "text")


# ---------------------------------------------------------------------------------------------
# reference material: the sheet's own Telugu column, shown read-only next to the editor
# ---------------------------------------------------------------------------------------------


def source_telugu() -> dict[str, str]:
    try:
        import openpyxl

        from tools.import_sheet import XLSX, clean, number, slug
    except ImportError:
        return {}
    if not XLSX.exists():
        return {}
    out = {}
    ws = openpyxl.load_workbook(XLSX, read_only=True)["Sheet1"]
    for row_no, row in enumerate(ws.iter_rows(values_only=True), 1):
        if row_no <= 3:
            continue
        no, te = number(row[0]), clean(row[4] if len(row) > 4 else "")
        if no and te:
            out.setdefault("v" + slug(no), te)
    return out


SOURCE_TE: dict[str, str] = {}


# ---------------------------------------------------------------------------------------------
# git, read-only
# ---------------------------------------------------------------------------------------------


def git(*args: str) -> str:
    try:
        r = subprocess.run(["git", *args], cwd=ROOT, capture_output=True, text=True, timeout=10)
        return r.stdout if r.returncode == 0 else ""
    except (OSError, subprocess.TimeoutExpired):
        return ""


def committed_doc():
    text = git("show", "HEAD:data/grantha.yaml")
    if not text:
        return None
    try:
        return store._yaml().load(StringIO(text))
    except Exception:
        return None


def committed_entries() -> dict[str, dict]:
    doc = committed_doc()
    if doc is None:
        return {}
    return {e["id"]: fields_of(e) for e in store.entries(doc)}


# ---------------------------------------------------------------------------------------------
# API
# ---------------------------------------------------------------------------------------------


def fields_of(e) -> dict[str, str]:
    """The editable fields of an entry. Every verse has a header slot, filled or not."""
    out = {k: str(e.get(k) or "") for k in FIELDS if k in e}
    if e["type"] == "verse":
        out = {"title": str(e.get("title") or ""), **out}
    return out


def flags() -> dict[str, str]:
    """Entries marked by hand to come back to: id -> why. They count as "has issues"."""
    if not FLAGS.exists():
        return {}
    with FLAGS.open("rb") as f:
        return {str(k): str(v) for k, v in tomllib.load(f).items()}


def lint(entry_id: str, text: str) -> list[str]:
    """HK.lint, minus the findings the author has looked at and accepted for this entry."""
    found = HK.lint(text)
    if not LINT_OK.exists():
        return found
    with LINT_OK.open("rb") as f:
        accepted = tomllib.load(f).get(entry_id) or []
    return [i for i in found if not any(i.startswith(a) for a in accepted)]


def label_of(e) -> str:
    body = str(e.get("title") or e.get("hk") or e.get("text") or "")
    return re.sub(r"\s+", " ", body)[:60]


def api_entries():
    doc = store.load()
    head = committed_entries()
    marked = flags()
    out = []
    group = "Preface (English)"
    for lang in ("en", "te"):
        group = "Preface (English)" if lang == "en" else "ముందుమాట (Telugu preface)"
        for b in doc["prefaces"][lang]["blocks"]:
            out.append(_row(b, group, head, marked))
    group = "Opening verses"
    for it in doc["items"]:
        if it["type"] == "section":
            group = re.sub(r"\s+", " ", str(it["hk"]))
        out.append(_row(it, group, head, marked))
    return {"entries": out, "git": git_summary()}


def _row(e, group, head, marked):
    issues = lint(e["id"], str(e["hk"])) if e.get("hk") else []
    if e.get("title"):
        issues += lint(e["id"], str(e["title"]))
    if e["type"] == "verse" and not str(e.get("en") or "").strip():
        issues.append("no English translation")
    if e["id"] in marked:
        issues.append(marked[e["id"]])
    current = fields_of(e)
    return {
        "id": e["id"],
        "type": e["type"],
        "no": e.get("no", ""),
        "label": label_of(e),
        "group": group,
        "issues": len(issues),
        "uncommitted": bool(head) and head.get(e["id"]) != current,
    }


def api_entry(entry_id: str):
    doc = store.load()
    e = store.find(doc, entry_id)
    head = committed_entries().get(entry_id)
    return {
        "id": e["id"],
        "type": e["type"],
        "no": e.get("no", ""),
        "fields": fields_of(e),
        "committed": head,
        "sourceTelugu": SOURCE_TE.get(entry_id, ""),
        "flag": flags().get(entry_id, ""),
    }


def api_preview(body):
    hk = str(body.get("hk", ""))
    script = body.get("script") or "telugu"
    if script not in render.config()["by_id"]:
        raise KeyError(script)
    with LOCK:
        spans = render.render_spans(hk, script, body.get("mode") or "zuddha")
        sarala_te = render.render(hk, "telugu", "saraLa")
    title = str(body.get("title", ""))
    with LOCK:
        title_shown = " ".join(render.render(title, script, body.get("mode") or "zuddha")) if title.strip() else ""
    return {
        "titleLine": title_shown,
        "titleZuddha": HK.to_zuddha(title),
        "titleIssues": [i for i in HK.lint(title) if not i.startswith("not in zuddha")],
        "lines": [" ".join(t for t, _, _ in line) for line in spans],
        "spans": spans,
        "saraLaTelugu": sarala_te,
        "zuddha": HK.to_zuddha(hk),
        "issues": lint(str(body.get("id", "")), hk),
    }


def api_preview_prose(body):
    """Prose with its inline ``$hk$`` runs rendered, for the box under a prose field."""
    script = body.get("script") or "telugu"
    if script not in render.config()["by_id"]:
        raise KeyError(script)
    mode = body.get("mode") or "zuddha"
    paras, issues = [], []
    with LOCK:
        for para in build.paragraphs(str(body.get("text", ""))):
            segs = []
            for kind, part in inline.split(para):
                if kind == "hk":
                    issues += [f"${part}$: {i}" for i in HK.lint(part)]
                    part = " ".join(render.render(part, script, mode))
                segs.append([kind, part])
            paras.append(segs)
    if str(body.get("text", "")).replace("\\$", "").count("$") % 2:
        issues.append("an inline $…$ run is not closed")
    return {"paragraphs": paras, "issues": issues}


def _apply(doc, entry_id: str, body) -> list[str]:
    """Write the changed fields of one entry into the in-memory doc; return which changed.
    In hk, | and || are first rewritten as . and .., and a verse is broken into one line per
    danda. A meaning in the wrong language is refused (ValueError) before anything is written."""
    e = store.find(doc, entry_id)
    changed = []
    for k in FIELDS:
        if k not in body or k not in fields_of(e):
            continue
        new = str(body[k]).replace("\r\n", "\n").strip("\n")
        if k == "hk":
            new = HK.dot_dandas(new)  # dandas are written . and .., never | and ||
            if e.get("type") == "verse":
                new = HK.break_lines(new)  # one line of the verse per line of text
        if str(e.get(k) or "") == new:
            continue
        if why := store.wrong_language(k, new):
            raise ValueError(f"{e.get('no') or entry_id}: {why}")
        if k == "title":
            store.set_title(e, new)
        else:
            e[k] = store.text(new)
        changed.append(k)
    return changed


def api_save(entry_id: str, body):
    with LOCK:
        doc = store.load()
        changed = _apply(doc, entry_id, body)
        if changed:
            store.save(doc)
            build.build(quiet=True)
    return {"saved": changed, **api_entry(entry_id), "git": git_summary()}


def api_save_all(body):
    """Save many entries at once: {"entries": {id: {field: value}}}. One write, one rebuild.
    Nothing is written if any id is unknown or any meaning is in the wrong language."""
    items = body.get("entries") or {}
    with LOCK:
        doc = store.load()
        saved = {eid: c for eid, vals in items.items() if (c := _apply(doc, eid, vals))}
        if saved:
            store.save(doc)
            build.build(quiet=True)
    return {"saved": saved, "git": git_summary()}


def api_add_block(body):
    """Append an empty block to the Telugu preface (the sheet had none)."""
    kind = "sloka" if body.get("type") == "sloka" else "prose"
    with LOCK:
        doc = store.load()
        blocks = doc["prefaces"]["te"]["blocks"]
        taken = {b["id"] for b in blocks}
        n = 1
        while f"pt{n:02d}" in taken:
            n += 1
        block = {"type": kind, "id": f"pt{n:02d}", ("hk" if kind == "sloka" else "text"): ""}
        blocks.append(block)
        store.save(doc)
    return {"id": block["id"]}


def git_summary():
    status = git("status", "--porcelain").strip()
    branch = git("rev-parse", "--abbrev-ref", "HEAD").strip()
    return {"branch": branch, "dirty": [ln for ln in status.split("\n") if ln], "repo": bool(branch)}


# ---------------------------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------------------------


class Handler(BaseHTTPRequestHandler):
    server_version = "krishnanandamayam"

    def log_message(self, fmt, *args):
        if "/api/preview" not in str(args[0] if args else ""):
            super().log_message(fmt, *args)

    def _json(self, payload, status=HTTPStatus.OK):
        data = json.dumps(payload, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def _data_error(self, err):
        """grantha.yaml cannot be read (merge conflict, bad YAML): say so, so the editor can show it."""
        kind = "conflict" if err.conflicts else "invalid"
        self._json({"error": str(err), "data_error": {"kind": kind, "conflicts": err.conflicts}}, HTTPStatus.CONFLICT)

    def _file(self, base, rel):
        path = (base / rel).resolve()
        if base.resolve() not in path.parents and path != base.resolve():
            return self.send_error(HTTPStatus.FORBIDDEN)
        if path.is_dir():
            path = path / "index.html"
        if not path.is_file():
            return self.send_error(HTTPStatus.NOT_FOUND)
        data = path.read_bytes()
        ctype = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        if ctype.startswith("text/") or ctype in ("application/json", "application/javascript"):
            ctype += "; charset=utf-8"
        self.send_response(HTTPStatus.OK)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        path = unquote(urlparse(self.path).path)
        try:
            if path == "/api/entries":
                return self._json(api_entries())
            if path.startswith("/api/entry/"):
                return self._json(api_entry(path.rsplit("/", 1)[1]))
            if path == "/api/scripts":
                return self._json({"scripts": [{"id": s["id"], "label": s["label"], "font": s["font"]} for s in render.config()["script"]]})
        except KeyError as err:
            return self._json({"error": f"not found: {err}"}, HTTPStatus.NOT_FOUND)
        except store.DataFileError as err:
            return self._data_error(err)
        if path in ("/edit", "/edit/"):
            return self._file(EDITOR, "editor.html")
        if path.startswith("/edit/"):
            return self._file(EDITOR, path[len("/edit/") :])
        return self._file(SITE, path.lstrip("/") or "index.html")

    def do_POST(self):
        # Only this machine's own pages may write: refuse cross-site requests.
        origin = self.headers.get("Origin")
        host = self.headers.get("Host", "")
        if origin and urlparse(origin).netloc != host:
            return self._json({"error": "cross-origin request refused"}, HTTPStatus.FORBIDDEN)
        if "application/json" not in (self.headers.get("Content-Type") or ""):
            return self._json({"error": "JSON only"}, HTTPStatus.UNSUPPORTED_MEDIA_TYPE)
        try:
            length = int(self.headers.get("Content-Length") or 0)
            body = json.loads(self.rfile.read(length) or b"{}")
            path = unquote(urlparse(self.path).path)
            if path == "/api/preview":
                return self._json(api_preview(body))
            if path == "/api/preview-prose":
                return self._json(api_preview_prose(body))
            if path == "/api/save-all":
                return self._json(api_save_all(body))
            if path.startswith("/api/entry/"):
                return self._json(api_save(path.rsplit("/", 1)[1], body))
            if path == "/api/preface/te/add":
                return self._json(api_add_block(body))
            return self._json({"error": "unknown endpoint"}, HTTPStatus.NOT_FOUND)
        except KeyError as err:
            return self._json({"error": f"not found: {err}"}, HTTPStatus.NOT_FOUND)
        except store.DataFileError as err:
            return self._data_error(err)
        except (ValueError, json.JSONDecodeError) as err:
            return self._json({"error": str(err)}, HTTPStatus.BAD_REQUEST)


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--port", type=int, default=8000)
    ap.add_argument("--host", default="127.0.0.1", help="0.0.0.0 to reach it from a phone on the same Wi-Fi")
    args = ap.parse_args()
    global SOURCE_TE
    SOURCE_TE = source_telugu()
    try:
        build.build()
    except store.DataFileError as err:
        print(f"warning: {err}; the editor will say so until it is resolved")
    httpd = ThreadingHTTPServer((args.host, args.port), Handler)
    print(f"reader  http://localhost:{args.port}/")
    print(f"editor  http://localhost:{args.port}/edit")
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print()


if __name__ == "__main__":
    main()
